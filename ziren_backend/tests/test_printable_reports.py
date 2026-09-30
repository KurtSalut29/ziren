"""
printable_reports — the two documents Reports & Export prints: the Incident
Records Report and the Narrative Reports bundle.

PDFs are built with page compression off (see _plain_pdf) so the text on the
page can be asserted directly from the bytes.

Run: pytest tests/test_printable_reports.py -v
"""

import csv
import io
from unittest.mock import MagicMock, patch

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient

from app.core.dependencies import get_current_user
from app.main import app
from app.services import printable_reports

client = TestClient(app)

AGENCY_ADMIN = {"id": "u1", "role": "agency_admin", "agency_id": "ag1", "agency_type": "BFP", "full_name": "Kurt Salut"}
PROVINCIAL_ADMIN = {"id": "u2", "role": "provincial_admin", "agency_type": "PNP", "full_name": "Prov Admin"}


@pytest.fixture(autouse=True)
def _plain_pdf(monkeypatch):
    from reportlab import rl_config
    monkeypatch.setattr(rl_config, "pageCompression", 0)
    yield
    app.dependency_overrides.clear()


def _text(pdf: bytes) -> str:
    """The text drawn on the pages, from an uncompressed PDF: every `(...) Tj`
    string, unescaped, one per line."""
    import re
    raw = pdf.decode("latin-1")
    parts = re.findall(r"\((.*?)(?<!\\)\) Tj", raw)
    unescape = {"\\(": "(", "\\)": ")", "\\267": "·", "\\227": "—"}
    out = []
    for p in parts:
        for k, v in unescape.items():
            p = p.replace(k, v)
        out.append(p)
    return "\n".join(out)


def _flat(text: str) -> str:
    """Text as read, not as laid out: a cell that wraps onto two lines still
    reads as one phrase."""
    return " ".join(text.split())


def _agency_db(name="BFP Naval Station", municipality="Naval", agency_type="BFP"):
    db = MagicMock()
    db.table.return_value.select.return_value.eq.return_value.maybe_single.return_value.execute.return_value = MagicMock(
        data={"name": name, "municipality": municipality, "agency_type": agency_type},
    )
    return db


def _incident(n, **kw):
    row = {
        "id": f"i{n}", "record_number": f"BFP-2026-{n:06d}", "status": "resolved", "severity": "critical",
        "incident_category": "fire", "location_address": f"Brgy. {n}, Naval, Biliran",
        "created_at": "2026-09-20T01:00:00Z", "dispatched_at": "2026-09-20T01:07:00Z",
        "resolved_at": "2026-09-20T02:00:00Z", "outcome": "fire_contained",
        "stations": {"name": "BFP Naval"}, "responder": {"full_name": "FO1 Cruz"},
        "report_text": f"Fire — Sunog sa bahay number {n}, may naiwan pang bata sa loob",
    }
    row.update(kw)
    return row


def _history(items, total=None, counts=None):
    return {
        "items": items,
        "total": len(items) if total is None else total,
        "counts": counts or {"total": len(items), "critical": 1, "resolved": 1, "cancelled": 0, "avg_response_minutes": 7.0},
    }


# ── Incident Records Report ────────────────────────────────────────────────

def test_incident_records_pdf_is_a_formatted_document():
    with patch("app.services.printable_reports.dispatch_service.get_incident_history",
               return_value=_history([_incident(1), _incident(2, severity="low", status="dispatched", resolved_at=None)])), \
         patch("app.services.printable_reports.get_supabase", return_value=_agency_db()):
        content, filename, media = printable_reports.build_incident_records(
            AGENCY_ADMIN, "pdf", start_date="2026-09-01", end_date="2026-09-30",
        )
    assert media == "application/pdf"
    assert filename.startswith("incident_records_") and filename.endswith(".pdf")
    assert content.startswith(b"%PDF")
    text = _flat(_text(content))
    for phrase in (
        "INCIDENT RECORDS REPORT", "Republic of the Philippines", "BUREAU OF FIRE PROTECTION",
        "BFP Naval Station, Naval", "September 1, 2026 to September 30, 2026", "BFP-2026-000001",
        "BFP-2026-000002", "Critical", "Low", "PREPARED BY", "NOTED BY", "Page 1 of 1",
        "Kurt Salut (BFP Agency Admin)", "Sep 20, 2026 9:00",
    ):
        assert phrase in text, phrase


def test_incident_records_pdf_says_so_when_nothing_matches():
    with patch("app.services.printable_reports.dispatch_service.get_incident_history",
               return_value=_history([], counts={"total": 0, "critical": 0, "resolved": 0, "cancelled": 0, "avg_response_minutes": None})), \
         patch("app.services.printable_reports.get_supabase", return_value=_agency_db()):
        content, _, _ = printable_reports.build_incident_records(AGENCY_ADMIN, "pdf")
    text = _text(content)
    assert "No incident matches this period and these filters." in text
    assert "All records to date" in text


def test_incident_records_csv_has_one_row_per_incident_and_passes_filters_through():
    with patch("app.services.printable_reports.dispatch_service.get_incident_history",
               return_value=_history([_incident(1)])) as hist, \
         patch("app.services.printable_reports.get_supabase", return_value=_agency_db()):
        content, filename, media = printable_reports.build_incident_records(
            AGENCY_ADMIN, "csv", start_date="2026-09-01", end_date="2026-09-30", status_filter="resolved", severity="critical",
        )
    assert media == "text/csv" and filename.endswith(".csv")
    rows = list(csv.reader(io.StringIO(content.decode("utf-8"))))
    assert rows[0] == printable_reports.EXPORT_HEADERS
    assert rows[0][3] == "Resident's Report"
    assert rows[1][0] == "BFP-2026-000001"
    assert rows[1][3] == "Fire — Sunog sa bahay number 1, may naiwan pang bata sa loob"
    assert rows[1][4] == "Critical"
    assert rows[1][9] == "7"  # minutes from report to dispatch
    kwargs = hist.call_args.kwargs
    assert (kwargs["status"], kwargs["severity"], kwargs["date_from"], kwargs["date_to"], kwargs["days"]) == (
        "resolved", "critical", "2026-09-01", "2026-09-30", 0,
    )
    # Scoping is the history query's own: the caller is passed straight through.
    assert hist.call_args.args[0] is AGENCY_ADMIN


def test_incident_records_pages_through_the_whole_window():
    page = printable_reports.dispatch_service.HISTORY_PAGE_MAX
    batches = [
        _history([_incident(i) for i in range(page)], total=page * 2 + 5),
        _history([_incident(i) for i in range(page)], total=page * 2 + 5),
        _history([_incident(i) for i in range(5)], total=page * 2 + 5),
    ]
    with patch("app.services.printable_reports.dispatch_service.get_incident_history", side_effect=batches) as hist, \
         patch("app.services.printable_reports.get_supabase", return_value=_agency_db()):
        content, _, _ = printable_reports.build_incident_records(AGENCY_ADMIN, "csv")
    assert [c.kwargs["offset"] for c in hist.call_args_list] == [0, page, page * 2]
    assert len(list(csv.reader(io.StringIO(content.decode("utf-8"))))) == page * 2 + 5 + 1


# ── What the resident said (2026-10-01) ────────────────────────────────────
#
# The records export said where, when, how bad and who went - but never what
# was reported. The resident's own words now go with every record.

@pytest.mark.parametrize("row, expected", [
    ({"report_text": "Fire — may sunog sa bahay"}, "Fire — may sunog sa bahay"),
    ({"report_text": "Fire — no additional details provided"}, "Fire — no additional details provided"),
    ({"report_text": "SOS — tulong po"}, "SOS — tulong po"),
    ({"report_text": "Fire — reported by voice recording — Sir tabang may kalayo"}, "(Voice) Sir tabang may kalayo"),
    ({"report_text": "Sunog — iniulat sa pamamagitan ng boses — may sunog po dito"}, "(Voice) may sunog po dito"),
    (
        {"report_text": "Fire — reported by voice recording — may kalayo sa balay",
         "signals": {"normalisation": {"text": "Fire — reported by voice recording — may kalayo sa balay namon"}}},
        "(Voice) may kalayo sa balay namon",
    ),
    ({"report_text": "Fire — reported by voice recording"}, "(Voice note — no transcript)"),
    ({"report_text": ""}, "—"),
])
def test_resident_words_strip_only_the_app_s_own_scaffolding(row, expected):
    assert printable_reports.resident_words(row) == expected


def test_resident_words_are_shortened_for_print_only_when_asked():
    long = {"report_text": "a" * 900}
    assert printable_reports.resident_words(long) == "a" * 900
    assert len(printable_reports.resident_words(long, limit=700)) == 700


def test_incident_records_pdf_has_a_resident_s_report_column():
    with patch("app.services.printable_reports.dispatch_service.get_incident_history",
               return_value=_history([_incident(1), _incident(2, report_text="Fire — reported by voice recording")])), \
         patch("app.services.printable_reports.get_supabase", return_value=_agency_db()):
        content, _, _ = printable_reports.build_incident_records(AGENCY_ADMIN, "pdf")
    text = _flat(_text(content))
    # A column header, not a label repeated under every record.
    assert text.count("Resident's Report") == 1
    assert "Resident's report:" not in text
    assert "Sunog sa bahay number 1, may naiwan pang bata sa loob" in text
    assert "(Voice note — no transcript)" in text
    # The note is the report's, not the resident's words: never quoted.
    assert "“" not in text


def test_incident_records_pdf_columns_fill_the_page_exactly():
    from reportlab.lib.units import inch
    from reportlab.platypus import Table

    captured = {}

    class Spy(Table):
        def __init__(self, data, colWidths=None, **kw):
            if data and len(data[0]) == 12:
                captured["w"] = colWidths
            super().__init__(data, colWidths=colWidths, **kw)

    # The builder imports Table inside the function, so patching the module
    # attribute reaches it.
    with patch("app.services.printable_reports.dispatch_service.get_incident_history",
               return_value=_history([_incident(1)])), \
         patch("app.services.printable_reports.get_supabase", return_value=_agency_db()), \
         patch("reportlab.platypus.Table", Spy):
        printable_reports.build_incident_records(AGENCY_ADMIN, "pdf")
    assert abs(sum(captured["w"]) - 10 * inch) < 0.01
    assert captured["w"][4] == max(captured["w"])  # the report column is the widest


def test_incident_records_xlsx_carries_the_report_wrapped():
    from openpyxl import load_workbook
    with patch("app.services.printable_reports.dispatch_service.get_incident_history", return_value=_history([_incident(1)])), \
         patch("app.services.printable_reports.get_supabase", return_value=_agency_db()):
        content, _, _ = printable_reports.build_incident_records(AGENCY_ADMIN, "xlsx")
    ws = load_workbook(io.BytesIO(content)).active
    assert ws["D1"].value == "Resident's Report"
    assert ws["D2"].value.endswith("may naiwan pang bata sa loob")
    assert ws["D2"].alignment.wrap_text is True
    assert ws.column_dimensions["D"].width == 60


def test_incident_records_xlsx_is_produced():
    with patch("app.services.printable_reports.dispatch_service.get_incident_history", return_value=_history([_incident(1)])), \
         patch("app.services.printable_reports.get_supabase", return_value=_agency_db()):
        content, filename, media = printable_reports.build_incident_records(AGENCY_ADMIN, "xlsx")
    assert content[:2] == b"PK" and filename.endswith(".xlsx")
    assert media.endswith("spreadsheetml.sheet")


@pytest.mark.parametrize("kwargs", [
    {"fmt": "docx"},
    {"fmt": "pdf", "status_filter": "everything"},
    {"fmt": "pdf", "severity": "apocalyptic"},
])
def test_incident_records_rejects_bad_parameters(kwargs):
    fmt = kwargs.pop("fmt")
    with pytest.raises(HTTPException) as exc:
        printable_reports.build_incident_records(AGENCY_ADMIN, fmt, **kwargs)
    assert exc.value.status_code == 422


def test_provincial_letterhead_names_the_agency_type_not_a_station():
    with patch("app.services.printable_reports.dispatch_service.get_incident_history", return_value=_history([_incident(1)])):
        content, _, _ = printable_reports.build_incident_records(PROVINCIAL_ADMIN, "pdf")
    assert b"PHILIPPINE NATIONAL POLICE" in content
    assert b"PNP \\227 all stations" in content or "PNP — all stations".encode("cp1252") in content


# ── Narrative Reports bundle ───────────────────────────────────────────────

def _loaded(n, status="finalized"):
    inc = {
        "id": f"i{n}", "record_number": f"PNP-2026-{n:06d}", "incident_category": "domestic_dispute_crime",
        "created_at": "2026-09-20T01:00:00Z", "location_address": "Brgy. Larrazabal, Naval",
        "stations": {"name": "PNP Naval", "agencies": {"agency_type": "PNP", "name": "PNP Naval Station", "municipality": "Naval"}},
        "users": {"full_name": "Juan Dela Cruz", "phone_number": "0917"},
    }
    rep = {"status": status, "reference_no": f"IRF-{n:04d}", "narrative": f"Narrative number {n}.", "details": {}}
    return inc, rep


def test_narrative_bundle_has_a_cover_page_and_every_form():
    def load(db, incident_id, actor):
        n = int(incident_id[1:])
        return _loaded(n, status="draft" if n == 2 else "finalized")

    with patch("app.services.printable_reports.incident_narrative_service.load_for_print", side_effect=load), \
         patch("app.services.printable_reports.get_supabase", return_value=_agency_db("PNP Naval Station", "Naval", "PNP")):
        content, filename, media = printable_reports.build_narrative_bundle(AGENCY_ADMIN, ["i1", "i2", "i1"])

    assert media == "application/pdf" and filename.startswith("narrative_reports_")
    text = _text(content)
    assert "NARRATIVE REPORTS" in text
    assert "2 (1 finalized, 1 draft)" in text          # the duplicate id is printed once
    assert text.count("INCIDENT RECORD FORM") == 2
    assert "IRF-0001" in text and "IRF-0002" in text
    assert "Narrative number 1." in text and "Narrative number 2." in text
    assert "Page 1 of" in text


def test_narrative_bundle_leaves_out_unwritten_reports_and_says_so():
    def load(db, incident_id, actor):
        if incident_id == "i2":
            raise HTTPException(status_code=404, detail="No narrative report")
        return _loaded(1)

    with patch("app.services.printable_reports.incident_narrative_service.load_for_print", side_effect=load), \
         patch("app.services.printable_reports.get_supabase", return_value=_agency_db()):
        content, _, _ = printable_reports.build_narrative_bundle(AGENCY_ADMIN, ["i1", "i2"])
    # The sentence wraps across lines on the page; join them back up.
    assert "1 selected incident has no narrative report written yet and is not included." in " ".join(_text(content).splitlines())


def test_narrative_bundle_refuses_a_report_outside_the_callers_scope():
    def load(db, incident_id, actor):
        raise HTTPException(status_code=403, detail="Not your agency")

    with patch("app.services.printable_reports.incident_narrative_service.load_for_print", side_effect=load), \
         patch("app.services.printable_reports.get_supabase", return_value=_agency_db()):
        with pytest.raises(HTTPException) as exc:
            printable_reports.build_narrative_bundle(AGENCY_ADMIN, ["i1"])
    assert exc.value.status_code == 403


def test_narrative_bundle_limits():
    with patch("app.services.printable_reports.get_supabase", return_value=_agency_db()):
        with pytest.raises(HTTPException) as empty:
            printable_reports.build_narrative_bundle(AGENCY_ADMIN, [" ", ""])
        with pytest.raises(HTTPException) as too_many:
            printable_reports.build_narrative_bundle(AGENCY_ADMIN, [f"i{n}" for n in range(51)])
    assert empty.value.status_code == 422
    assert too_many.value.status_code == 422


# ── Routes ─────────────────────────────────────────────────────────────────

def test_incident_records_route_is_not_swallowed_by_the_report_type_route():
    app.dependency_overrides[get_current_user] = lambda: AGENCY_ADMIN
    with patch("app.routers.reports.printable_reports.build_incident_records",
               return_value=(b"%PDF-1.4", "incident_records_x.pdf", "application/pdf")) as fn:
        res = client.get("/reports/incident_records?format=pdf&status=open&severity=high&start_date=2026-09-01")
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf"
    assert res.headers["content-disposition"] == 'inline; filename="incident_records_x.pdf"'
    assert fn.call_args.kwargs == {"start_date": "2026-09-01", "end_date": None, "status_filter": "open", "severity": "high"}


def test_narrative_reports_route_splits_ids():
    app.dependency_overrides[get_current_user] = lambda: PROVINCIAL_ADMIN
    with patch("app.routers.reports.printable_reports.build_narrative_bundle",
               return_value=(b"%PDF-1.4", "narrative_reports_x.pdf", "application/pdf")) as fn:
        res = client.get("/reports/narrative_reports?ids=a,b,c")
    assert res.status_code == 200
    assert fn.call_args.args == (PROVINCIAL_ADMIN, ["a", "b", "c"])


def test_printable_routes_are_admin_only():
    app.dependency_overrides[get_current_user] = lambda: {"id": "r1", "role": "responder", "agency_id": "ag1"}
    assert client.get("/reports/incident_records").status_code == 403
    assert client.get("/reports/narrative_reports?ids=a").status_code == 403
